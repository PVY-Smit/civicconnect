"""Export the Access Matrix and Status Model sheets of the registers workbook as the test fixtures.

    python tools/export_register_fixtures.py          # rewrite tests/fixtures/*.json
    python tools/export_register_fixtures.py --check  # exit 1 if either fixture differs from the workbook

The policy (ADR-006) and transition table (ADR-005) are tested against these fixtures, so every cell of
the controlled registers is a test case. RSK-16 is the risk of the fixtures and the registers drifting
apart; this script is the one way the fixtures are produced, and --check shows whether they still match.
The source line records the workbook's SHA-256, so a fixture can be traced to the exact file it came from.

Needs openpyxl (pip install openpyxl). Only the test tooling uses it; the application does not.
"""
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "docs" / "requirements" / "CivicConnect_M1_Registers.xlsx"
FIXTURES = ROOT / "tests" / "fixtures"
ROLES = ["Requester", "Staff", "Coordinator", "Manager"]


def source(sheet):
    digest = hashlib.sha256(WORKBOOK.read_bytes()).hexdigest()[:12]
    return f"{sheet} sheet, docs/requirements/{WORKBOOK.name} (sha256 {digest})"


def access_matrix(wb):
    ws = wb["Access Matrix"]
    header = [c.value for c in ws[2]]
    assert header[:5] == ["Function", *ROLES], header
    functions = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row[0]:
            continue
        cells = row[1:5]
        assert all(v in ("Yes", "No") for v in cells), (row[0], cells)
        functions.append({"function": row[0], **{role: v == "Yes" for role, v in zip(ROLES, cells)}})
    return {"source": source("Access Matrix"), "roles": ROLES, "functions": functions}


def status_model(wb):
    ws = wb["Status Model"]
    transitions = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row[0]:
            continue
        frm, to, roles, guard = row[:4]
        transitions.append({"from": frm, "to": to, "roles": [r.strip() for r in str(roles).split(",")], "guard": guard})
    return {"source": source("Status Model"), "transitions": transitions}


def main():
    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    outputs = {"access-matrix.json": access_matrix(wb), "status-model.json": status_model(wb)}
    check = "--check" in sys.argv
    stale = []
    for name, data in outputs.items():
        text = json.dumps(data, indent=2, ensure_ascii=False) + "\n"
        path = FIXTURES / name
        if check:
            if not path.exists() or path.read_text(encoding="utf-8") != text:
                stale.append(name)
        else:
            path.write_text(text, encoding="utf-8", newline="\n")
            print(f"wrote {path.relative_to(ROOT)}")
    if check:
        if stale:
            print("out of date with the workbook: " + ", ".join(stale))
            sys.exit(1)
        print("fixtures match the workbook")


if __name__ == "__main__":
    main()
